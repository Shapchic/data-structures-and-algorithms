#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Скрипт генерации детального руководства к защите ЛР 3 в формате DOCX."""

import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Установка фонового цвета ячейки таблицы."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Установка внутренних отступов ячейки."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Установка тонких аккуратных границ таблицы."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, question, answer):
    """Добавление стилизованной карточки Вопрос-Ответ с левой цветной полосой."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Левая цветная полоса
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_q_title = p.add_run("❓ Каверзный вопрос преподавателя:\n")
    run_q_title.bold = True
    run_q_title.font.size = Pt(11)
    run_q_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_q = p.add_run(f"«{question}»\n\n")
    run_q.italic = True
    run_q.font.size = Pt(10.5)
    run_q.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    
    run_a_title = p.add_run("💡 Безупречный ответ на защите:\n")
    run_a_title.bold = True
    run_a_title.font.size = Pt(11)
    run_a_title.font.color.rgb = RGBColor(0x19, 0x6F, 0x3D)
    
    run_a = p.add_run(answer)
    run_a.font.size = Pt(10.5)
    run_a.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    
    # Пустой абзац после карточки
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    """Добавление моноширинного блока кода с фоном."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="5B9BD5"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_defense_docx():
    doc = Document()
    
    # Настройка полей страницы
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # ------------------ ТИТУЛЬНЫЙ БЛОК ------------------
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(2)
    r_disc = p_meta.add_run("СТРУКТУРЫ И АЛГОРИТМЫ ОБРАБОТКИ ДАННЫХ • МОДУЛЬ 2\n")
    r_disc.font.size = Pt(9.5)
    r_disc.font.bold = True
    r_disc.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Полное руководство к защите ЛР № 3\n«Простые сортировки и QuickSort с рандомизацией»")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Вариант 21 (seed = 51) • Шпаргалка ответов на самые каверзные вопросы")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    
    doc.add_heading("Глава 1. Фундаментальный теоретический базис", level=1)
    
    p = doc.add_paragraph(
        "На защите преподаватель проверяет не просто способность запустить код, "
        "а понимание глубинных математических ограничений вычислительных моделей, "
        "поведение алгоритмов при различных граничных условиях и умение обосновывать "
        "наблюдаемые аномалии в замерах."
    )
    p.paragraph_format.space_after = Pt(8)
    
    doc.add_heading("1.1. Информационно-теоретическая нижняя граница сложности", level=2)
    
    add_callout(
        doc,
        "Почему ни один алгоритм сортировки сравнениями не может работать быстрее Ω(n log n) "
        "в худшем и среднем случаях? Можно ли отсортировать быстрее?",
        "1. Любой алгоритм сортировки сравнениями моделируется бинарным деревом решений (Decision Tree).\n"
        "2. Для массива длины n существует ровно n! возможных перестановок, каждая из которых "
        "обязана быть представлена хотя бы в одном листе дерева решений, иначе алгоритм на каком-то входе ошибётся. "
        "Число листьев L >= n!.\n"
        "3. Бинарное дерево высоты h может иметь максимум 2^h листьев. Отсюда 2^h >= n!, "
        "что даёт h >= log2(n!).\n"
        "4. По формуле Стирлинга log2(n!) = n log2 n - n log2 e + O(log n) = Ω(n log n).\n"
        "Высота дерева h — это число сравнений в худшем случае. Следовательно, худший случай любого "
        "алгоритма на сравнениях есть Ω(n log n).\n"
        "5. Обойти эту границу можно ТОЛЬКО выйдя за рамки модели сравнений (например, Counting Sort, "
        "Radix Sort за O(n * k), используя побитовую структуру ключа и адресацию в память)."
    )
    
    doc.add_heading("1.2. Свойства: Стабильность, In-Place и Адаптивность", level=2)
    
    p_props = doc.add_paragraph(
        "• Стабильность (Stability): алгоритм сохраняет относительный порядок элементов с одинаковыми ключами. "
        "Критично при многокритериальной сортировке таблиц (сначала по имени, затем по отделу).\n"
        "• Сортировка на месте (In-Place): алгоритм использует O(1) вспомогательной памяти сверх исходного массива "
        "(для QuickSort допускается O(log n) для стека подзадач).\n"
        "• Адаптивность (Adaptivity): время работы алгоритма уменьшается при уменьшении числа инверсий во входе "
        "(например, вставки работают за O(n) на почти упорядоченном массиве)."
    )
    p_props.paragraph_format.space_after = Pt(10)
    
    # ------------------ ГЛАВА 2 ------------------
    doc.add_heading("Глава 2. Анализ четырёх алгоритмов и инварианты", level=1)
    
    doc.add_heading("2.1. Сортировка пузырьком (Bubble Sort)", level=2)
    p_bs = doc.add_paragraph(
        "Идея: попарные сравнения смежных элементов res[j] и res[j+1] с всплытием максимума в конец.\n"
        "Инвариант: к началу i-й итерации подмассив res[n-i .. n-1] отсортирован и содержит i наибольших элементов входа.\n"
        "Оптимизация: флаг swapped прерывает алгоритм, если за проход не было обменов."
    )
    p_bs.paragraph_format.space_after = Pt(6)
    
    add_callout(
        doc,
        "Сколько сравнений делает пузырёк с досрочным выходом на упорядоченном массиве и почему?",
        "Ровно n - 1 сравнений. Внешний цикл выполняет только 1 итерацию (i = 0). "
        "Внутренний цикл проверяет пары (0, 1), (1, 2), ..., (n-2, n-1) — ровно n - 1 проверок. "
        "Ни одного обмена не происходит, флаг swapped остаётся False, срабатывает break. Сложность Θ(n)."
    )
    
    doc.add_heading("2.2. Сортировка вставками (Insertion Sort)", level=2)
    p_is = doc.add_paragraph(
        "Идея: расширение отсортированного префикса res[0 .. i-1] путём вставки элемента res[i] на своё место.\n"
        "Инвариант: подмассив res[0 .. i-1] всегда отсортирован.\n"
        "Связь с инверсиями: число сдвигов строго равно числу инверсий во входном массиве: M(n) = I. "
        "Число сравнений C(n) <= I + n - 1. Сложность Θ(n + I)."
    )
    p_is.paragraph_format.space_after = Pt(6)
    
    add_callout(
        doc,
        "Почему Insertion Sort работает в 2.3 раза быстрее Bubble Sort на случайных данных, если у обоих O(n^2)?",
        "1. Перемещения: Вставка выполняет сдвиг res[j+1] = res[j] (1 запись в память), удерживая вставляемый "
        "элемент в локальной переменной. Пузырёк выполняет кортежный обмен (2 записи в память + оверхед кортежа Python).\n"
        "2. Сравнения: Вставка на случайных данных делает в среднем n^2 / 4 сравнений (останавливается в середине префикса), "
        "а пузырёк сравнивает почти всю неотсортированную зону (~ n^2 / 2).\n"
        "3. Локальность кэша: Сдвиги элементов идут в непрерывном массиве L1-кэша процессора."
    )
    
    doc.add_heading("2.3. Сортировка выбором (Selection Sort)", level=2)
    p_ss = doc.add_paragraph(
        "Идея: поиск индекса минимального элемента в суффиксе res[i .. n-1] и обмен с res[i].\n"
        "Инвариант: префикс res[0 .. i-1] содержит i наименьших элементов массива в отсортированном порядке."
    )
    p_ss.paragraph_format.space_after = Pt(6)
    
    add_callout(
        doc,
        "Почему сортировка выбором делает строго n(n-1)/2 сравнений абсолютно на любом входе?",
        "Потому что в алгоритме отсутствует возможность ранней остановки: чтобы доказать, что элемент res[min_idx] "
        "действительно минимален среди оставшихся n - 1 - i элементов, внутренний цикл обязан проверить каждый из них "
        "до самого конца. Число сравнений: sum_{i=0}^{n-2} (n - 1 - i) = n(n - 1)/2 во всех случаях (лучший, средний, худший)."
    )
    
    doc.add_heading("2.4. Быстрая сортировка (QuickSort) и 3-Way Partitioning", level=2)
    p_qs = doc.add_paragraph(
        "Идея: выбор опорного элемента, трёхчастное разбиение Дейкстры на зоны [< pivot, == pivot, > pivot] "
        "и последующая сортировка левой и правой зон через стек."
    )
    p_qs.paragraph_format.space_after = Pt(6)
    
    add_callout(
        doc,
        "В чём преимущество трёхчастного разбиения Дейкстры перед разбиением Ломуто на массиве из одинаковых элементов?",
        "Схема Ломуто на массиве из одинаковых элементов [5, 5, ..., 5] деградирует до худшего случая O(n^2), "
        "так как все равные элементы попадают в одну половину.\n"
        "Трёхчастное разбиение Дейкстры за 1 проход собирает все одинаковые элементы в центральной зоне [lt .. gt] "
        "и исключает их из дальнейшей сортировки. Алгоритм завершается за линейное время Θ(n)!"
    )
    
    add_callout(
        doc,
        "Как вы предотвратили переполнение стека (RecursionError) в QuickSort при больших n?",
        "1. Перешли от рекурсии к итеративному циклу со стеком границ отрезков.\n"
        "2. Применили оптимизацию хвостовой рекурсии: в стек сначала помещается БОЛЬШИЙ подмассив, "
        "а МЕНЬШИЙ обрабатывается первым.\n"
        "Так как меньший подмассив имеет размер <= n/2, максимальная глубина стека в памяти гарантированно "
        "не превышает ceil(log2 n). Для n = 100 000 глубина стека составляет не более 17 элементов!"
    )
    
    # ------------------ ГЛАВА 3 ------------------
    doc.add_heading("Глава 3. Демонстрация стабильности", level=1)
    
    p_stab = doc.add_paragraph(
        "Для исследования стабильности использовался массив пар (ключ, метка) с повторяющимися ключами: "
        "[(2, 'a'), (1, 'b'), (2, 'c'), (1, 'd')]. Сортировка выполнялась строго по первому элементу пары."
    )
    p_stab.paragraph_format.space_after = Pt(6)
    
    # Таблица стабильности
    tbl_stab = doc.add_table(rows=5, cols=4)
    tbl_stab.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_stab)
    
    headers = ["Алгоритм", "Результат сортировки", "Порядок меток для 2", "Вердикт"]
    for j, h in enumerate(headers):
        cell = tbl_stab.cell(0, j)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    data_stab = [
        ("Bubble Sort", "[(1, 'b'), (1, 'd'), (2, 'a'), (2, 'c')]", "['a', 'c'] (исходный)", "СТАБИЛЬНА"),
        ("Insertion Sort", "[(1, 'b'), (1, 'd'), (2, 'a'), (2, 'c')]", "['a', 'c'] (исходный)", "СТАБИЛЬНА"),
        ("Selection Sort", "[(1, 'b'), (1, 'd'), (2, 'c'), (2, 'a')]", "['c', 'a'] (инвертирован!)", "НЕ СТАБИЛЬНА"),
        ("QuickSort", "[(1, 'd'), (1, 'b'), (2, 'c'), (2, 'a')]", "['c', 'a'] (инвертирован!)", "НЕ СТАБИЛЬНА"),
    ]
    
    for i, row in enumerate(data_stab):
        bg = "FFFFFF" if i % 2 == 0 else "F9FAFB"
        for j, val in enumerate(row):
            cell = tbl_stab.cell(i + 1, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 3:
                r.bold = True
                r.font.color.rgb = RGBColor(0x19, 0x6F, 0x3D) if "СТАБИЛЬНА" == val else RGBColor(0xC0, 0x39, 0x2B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    add_callout(
        doc,
        "Объясните на пальцах, почему Selection Sort потерял стабильность на парах?",
        "На первом шаге алгоритм ищет минимум среди [(2, 'a'), (1, 'b'), (2, 'c'), (1, 'd')]. "
        "Минимум — это (1, 'b') на позиции 1. Алгоритм меняет местами элемент (2, 'a') на позиции 0 и (1, 'b'). "
        "Элемент (2, 'a') улетает на позицию 1. На следующем шаге ищется минимум суффикса — это (1, 'd') на позиции 3. "
        "Алгоритм меняет местами позицию 1 (где теперь стоит (2, 'a')) и позицию 3. В результате (2, 'a') оказывается "
        "в самом конце массива — ПОСЛЕ (2, 'c')! Относительный порядок одинаковых ключей 2 нарушен."
    )
    
    # ------------------ ГЛАВА 4 ------------------
    doc.add_heading("Глава 4. Результаты бенчмарков варианта 21 (seed = 51)", level=1)
    
    p_bch = doc.add_paragraph(
        "Замеры проводились на CPU Intel Core i5-11400F в Python 3.14.7 по схеме: 1 прогрев + 5 повторов (медиана)."
    )
    p_bch.paragraph_format.space_after = Pt(6)
    
    # Таблица наклонов
    tbl_slopes = doc.add_table(rows=5, cols=5)
    tbl_slopes.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_slopes)
    
    sl_headers = ["Алгоритм", "Упорядоченный вход", "Случайный вход", "Обратный вход", "Теория"]
    for j, h in enumerate(sl_headers):
        cell = tbl_slopes.cell(0, j)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    data_slopes = [
        ("bubble_sort", "1.042 (Θ(n))", "2.064 (Θ(n²))", "2.062 (Θ(n²))", "Лучший 1.0, худший 2.0"),
        ("insertion_sort", "1.050 (Θ(n))", "2.046 (Θ(n²))", "2.035 (Θ(n²))", "Лучший 1.0, худший 2.0"),
        ("selection_sort", "2.030 (Θ(n²))", "2.032 (Θ(n²))", "2.035 (Θ(n²))", "Всегда строго 2.0"),
        ("quick_sort", "1.132 (O(n log n))", "1.114 (O(n log n))", "1.130 (O(n log n))", "1 + 1/ln(n) ≈ 1.12"),
    ]
    
    for i, row in enumerate(data_slopes):
        bg = "FFFFFF" if i % 2 == 0 else "F9FAFB"
        for j, val in enumerate(row):
            cell = tbl_slopes.cell(i + 1, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j in (1, 2, 3):
                r.bold = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # Вставка графиков
    img_dir = Path("reports/images")
    if (img_dir / "lab03_sorts_loglog.png").exists():
        doc.add_heading("4.1. Логарифмические графики времени работы", level=2)
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(str(img_dir / "lab03_sorts_loglog.png"), width=Inches(6.2))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap1 = p_cap1.add_run("Рис. 1. Сравнение простых сортировок и QuickSort в осях log-log")
        r_cap1.font.size = Pt(9.0)
        r_cap1.font.italic = True
    
    if (img_dir / "lab03_quicksort.png").exists():
        doc.add_heading("4.2. Масштабирование QuickSort до 100 000 элементов", level=2)
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(str(img_dir / "lab03_quicksort.png"), width=Inches(5.0))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = p_cap2.add_run("Рис. 2. Масштабирование QuickSort с рандомизацией на трёх классах входов")
        r_cap2.font.size = Pt(9.0)
        r_cap2.font.italic = True
    
    # ------------------ ГЛАВА 5 ------------------
    doc.add_heading("Глава 5. Топ-10 каверзных вопросов преподавателя", level=1)
    
    add_callout(
        doc,
        "Вопрос 1: Зачем при замере бенчмарка передавать свежую копию list(data) на каждый повтор?",
        "Все сортировки сортируют список на месте. Если не передать свежую копию, то первый повтор отсортирует массив, "
        "а 2, 3, 4 и 5-й повторы будут замерять время на УЖЕ ОТСОРТИРОВАННОМ массиве! В результате вместо худшего случая "
        "обратного массива вставки покажут лучший случай O(n). Свежая копия гарантирует воспроизводимость и чистоту замеров."
    )
    
    add_callout(
        doc,
        "Вопрос 2: Почему наклон QuickSort на log-log графике равен 1.11...1.13, а не строго 1.00?",
        "Для степенной зависимости t = C * n^k наклон равен k. Для t = C * n ln n наклон равен производной "
        "d(log t)/d(log n) = 1 + 1/ln(n). На размерах n от 1 000 до 100 000 величина 1/ln(n) меняется от 0.14 до 0.09. "
        "Среднее значение равно как раз 1.12. Это строгое математическое подтверждение сложности O(n log n)!"
    )
    
    add_callout(
        doc,
        "Вопрос 3: Почему вычисление медианы лучше среднего арифметического при бенчмаркинге?",
        "Медиана обладает 50% точкой срыва (breakdown point) и нечувствительна к резким аномальным выбросам "
        "(outliers), вызванным переключением контекста планировщиком Windows, системными прерываниями и работой сборщика мусора Python."
    )
    
    add_callout(
        doc,
        "Вопрос 4: Можно ли ускорить сортировку вставками бинарным поиском до O(n log n)?",
        "Нет. Бинарный поиск сократит число сравнений до O(n log n), но число физических сдвигов элементов "
        "в памяти останется равным числу инверсий Θ(n^2). Сдвиг в массиве требует O(n) времени, поэтому итоговая сложность останется Θ(n^2)."
    )
    
    add_callout(
        doc,
        "Вопрос 5: Что такое «проблема черепах и зайцев» в Bubble Sort и как её решить?",
        "Большие элементы в начале («зайцы») перемещаются в конец за 1 проход. Маленькие элементы в конце («черепахи») "
        "смещаются влево всего на 1 позицию за проход. Если массив почти отсортирован, но наименьший элемент стоит в самом конце, "
        "пузырьку потребуется n - 1 проходов. Проблема решается шейкерной сортировкой (Cocktail Shaker Sort), обходящей массив в обоих направлениях."
    )
    
    add_callout(
        doc,
        "Вопрос 6: Почему в Selection Sort число сравнений на обратном массиве и упорядоченном одинаково, а число обменов разное?",
        "Сравнения ищут минимум безусловным сканированием суффикса (всегда n(n-1)/2). Обмен res[i], res[min_idx] выполняется, "
        "только если min_idx != i. На упорядоченном массиве всегда min_idx == i (0 обменов). На обратном массиве меняются "
        "симметричные элементы с концов к центру — ровно floor(n / 2) обменов (на n = 8000: ровно 4000 обменов)."
    )
    
    add_callout(
        doc,
        "Вопрос 7: Какие два инварианта сортировки обязательно проверять и почему одного недостаточно?",
        "1. Упорядоченность выхода: all(res[i] <= res[i+1]).\n"
        "2. Сохранение мультимножества: Counter(res) == Counter(a).\n"
        "Одной упорядоченности недостаточно, так как функция, возвращающая [] или [0] * n, отсортирована, но неверна. "
        "Одного равенства мультимножеств недостаточно, так как возврат исходного неотсортированного массива сохраняет мультимножество, "
        "но не сортирует. Корректность — это конъюнкция обоих свойств."
    )
    
    add_callout(
        doc,
        "Вопрос 8: Где в реальных промышленных библиотеках применяется Insertion Sort?",
        "Сортировка вставками лежит в основе стандартных алгоритмов Timsort (Python sorted, Java Arrays.sort) и Introsort (C++ std::sort). "
        "Когда в ходе рекурсивного деления QuickSort или MergeSort размер подмассива становится меньше порогового (обычно 16–32 элемента), "
        "алгоритм переключается на Insertion Sort, так как на малых n её минимальные накладные расходы и отсутствие ветвлений обгоняют любые O(n log n) алгоритмы."
    )
    
    add_callout(
        doc,
        "Вопрос 9: Какова вероятность худшего случая QuickSort при случайном выборе опорного элемента?",
        "Чтобы алгоритм работал за O(n^2), случайный генератор должен на каждом шаге выбирать минимальный или максимальный элемент. "
        "Вероятность этого равна 2^(n-1) / n!. Для n = 100 эта вероятность меньше 10^(-128). "
        "На практике возникновение худшего случая при честной рандомизации невозможно."
    )
    
    add_callout(
        doc,
        "Вопрос 10: Чем в памяти Python отличается список целых чисел от C-массива int?",
        "В Си массив int — это непрерывный кусок сырой памяти из 4-байтных чисел. "
        "В Python список list — это массив указателей (PyObject*), где каждый элемент — отдельный объект на куче "
        "с заголовком (счётчик ссылок ob_refcnt, указатель на тип ob_type, размер ob_size и само значение). "
        "Поэтому обмен a[i], a[j] в Python меняет местами 8-байтные указатели, не перемещая сами объекты чисел в куче."
    )
    
    out_docx = Path("reports/lab03-defense-guide.docx")
    doc.save(str(out_docx))
    print(f"DOCX файл успешно создан: {out_docx}")

if __name__ == "__main__":
    build_defense_docx()
