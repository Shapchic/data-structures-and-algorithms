#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Скрипт генерации детального руководства к защите ЛР 4 в формате DOCX."""

import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Установка фонового цвета ячейки таблицы."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Установка внутренних отступов ячейки."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Установка тонких аккуратных границ таблицы."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, question, answer):
    """Добавление выделенного визуального блока с вопросом и ответом."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F497D"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_q = p.add_run(question)
    run_q.bold = True
    run_q.font.size = Pt(10.5)
    run_q.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    run_a = p2.add_run(answer)
    run_a.font.size = Pt(10)
    run_a.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_defense_docx():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = False
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Лабораторная работа № 4 | Руководство к защите")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Кафедра программной инженерии — САОД 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("Руководство по защите лабораторной работы № 4")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("«MergeSort, Counting, Radix и HeapSort: оценка производительности»\nВариант 21 (seed = 51)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading("1. Сводная классификационная таблица шести алгоритмов", level=1)
    
    headers = ["Алгоритм", "Идея", "Лучшее", "Среднее", "Худшее", "Память", "Стабилен?", "In-place?"]
    table = doc.add_table(rows=7, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    widths = [Inches(1.1), Inches(1.2), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.7), Inches(0.6), Inches(0.5)]
    
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.width = widths[col_idx]
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    data_rows = [
        ("MergeSort", "Разделяй и слияние", "O(n log n)", "O(n log n)", "O(n log n)", "O(n)", "ДА", "НЕТ"),
        ("CountingSort", "Индексация ключами", "O(n + k)", "O(n + k)", "O(n + k)", "O(n + k)", "ДА", "НЕТ"),
        ("RadixSort (LSD)", "Поразрядно от мл.", "O(d·(n+b))", "O(d·(n+b))", "O(d·(n+b))", "O(n+b)", "ДА", "НЕТ"),
        ("HeapSort", "Двоичная max-куча", "O(n log n)", "O(n log n)", "O(n log n)", "O(1)", "НЕТ", "ДА"),
        ("QuickSort (ЛР 3)", "3-way partition", "O(n log n)", "O(n log n)", "O(n^2)*", "O(log n)", "НЕТ", "ДА"),
        ("sorted() Timsort", "Адаптивный гибрид", "O(n)", "O(n log n)", "O(n log n)", "O(n)", "ДА", "НЕТ"),
    ]
    
    for row_idx, row_data in enumerate(data_rows, start=1):
        bg = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            if col_idx in (2, 3, 4, 5, 6, 7):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    doc.add_heading("2. Ключевые вопросы для защиты (Банк ответов)", level=1)
    
    add_callout(
        doc,
        "Вопрос 1: Докажите нижнюю информационно-теоретическую границу Omega(n log n) для сортировок сравнением.",
        "Модель дерева решений: любой алгоритм парных сравнений образует бинарное дерево решений, где листья — возможные n! "
        "перестановок массива длины n. Число листьев L >= n!. Бинарное дерево высоты h содержит не более 2^h листьев: "
        "2^h >= L >= n!, следовательно h >= log2(n!). По приближению Стирлинга log2(n!) = Omega(n log n). "
        "Так как высота h соответствует наибольшему числу сравнений в худшем случае, сложность ограничена снизу Omega(n log n)."
    )

    add_callout(
        doc,
        "Вопрос 2: За счёт чего CountingSort и RadixSort работают быстрее Omega(n log n)?",
        "Они не сравнивают элементы друг с другом! Они используют цифровые (адресные) свойства ключей: "
        "значение ключа напрямую используется как индекс ячейки массива за O(1). "
        "Поскольку они выходят за рамки модели дерева решений, нижняя граница сравнений к ним неприменима."
    )

    add_callout(
        doc,
        "Вопрос 3: Почему в CountingSort обход входного массива в фазе 3 идёт строго с конца (reversed)?",
        "Это строгое условие стабильности! При обходе с конца элемент с ключом x, стоявший во входе позже (правее), "
        "обработается первым и займёт наибольшую доступную ячейку count[x] - 1. Предшествующий ему равный элемент встанет на позицию "
        "count[x] - 2. Исходный относительный порядок одинаковых ключей полностью сохраняется."
    )

    add_callout(
        doc,
        "Вопрос 4: Почему в RadixSort LSD поразрядный проход ОБЯЗАН быть устойчивым (стабильным)?",
        "Перед обработкой j-го разряда массив уже отсортирован по младшим j-1 разрядам. Если у двух чисел одинаковая цифра "
        "в текущем j-м разряде, их правильный порядок определяется именно младшими разрядами. "
        "Если процедура подсчёта не сохранит этот порядок, результат сортировки по всем предыдущим разрядам будет полностью уничтожен!"
    )

    add_callout(
        doc,
        "Вопрос 5: Докажите, почему фаза построения кучи (Heapify) в HeapSort работает за O(n), а не за O(n log n).",
        "Мы строим кучу снизу вверх от (n-2)//2 до 0. Большинство узлов находятся на нижних уровнях кучи и просеиваются на малую высоту: "
        "n/2 листьев не просеиваются вовсе, n/4 узлов имеют высоту спуска 1, n/8 узлов — высоту 2. "
        "Суммарное число операций: S = sum_{h=0}^{log n} (n / 2^(h+1)) * h = (n/2) * sum_{h=0}^inf (h / 2^h). "
        "Сумма ряда sum (h / 2^h) равна ровно 2. Получаем S = (n/2) * 2 = n = O(n)."
    )

    add_callout(
        doc,
        "Вопрос 6: Почему HeapSort на практике работает медленнее QuickSort при одинаковой сложности O(n log n)?",
        "Из-за плохой локальности данных в памяти (L1/L2 Cache Misses). В QuickSort массив обходится последовательно, "
        "что отлично кэшируется процессором. В HeapSort процедура sift_down прыгает по индексам 2*i + 1, постоянно обращаясь "
        "к разным нелокальным областям памяти и вызывая кэш-промахи."
    )

    add_callout(
        doc,
        "Вопрос 7: Почему CountingSort неприменим на распределении random_wide (диапазон ключей k = 10^9)?",
        "CountingSort требует массив счётчиков размера k = key_max + 1. Для k = 10^9 список из 10^9 целых чисел в Python "
        "потребует около 8 ГБ оперативной памяти и 10^9 операций инициализации. При n = 100 000 это приводит к исчерпанию ресурсов (MemoryError)."
    )

    add_callout(
        doc,
        "Вопрос 8: Почему в MergeSort условие слияния обязательно должно быть '<=', а не '<'?",
        "Условие left[i] <= right[j] гарантирует, что при равенстве ключей элемент из левой половины берётся первым. "
        "Так как в исходном массиве элементы левой половины предшествовали элементам правой, их взаимный порядок сохраняется. "
        "Если поставить строгое '<', при равенстве раньше попадёт элемент из правой половины, и сортировка станет нестабильной."
    )

    add_callout(
        doc,
        "Вопрос 9: Почему HeapSort является строго in-place алгоритмом с памятью O(1)?",
        "Потому что двоичное дерево кучи хранится прямо в исходном массиве без создания узлов-объектов и указателей. "
        "Связи между родителями и потомками вычисляются чисто арифметически по индексам (2*i + 1 и 2*i + 2). "
        "Процедура sift_down реализована итеративно циклом while, не расходуя стек вызовов."
    )

    add_callout(
        doc,
        "Вопрос 10: Чем Timsort (sorted()) отличается от классического MergeSort?",
        "Timsort — адаптивный гибрид. Он ищет в массиве уже упорядоченные участки (run). На отсортированных данных Timsort "
        "делает один линейный проход за Theta(n). Короткие участки упорядочиваются бинарными вставками (Insertion Sort), "
        "а слияние длинных серий оптимизировано режимом «галопа» (galloping)."
    )

    out_docx = Path("reports/lab04-defense-guide.docx")
    doc.save(str(out_docx))
    print(f"DOCX файл успешно создан: {out_docx}")

if __name__ == "__main__":
    build_defense_docx()
